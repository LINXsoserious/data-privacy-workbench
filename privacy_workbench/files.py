"""Bounded table loading. Uploaded data is never written to disk."""

import csv
from datetime import date, datetime
import io
from pathlib import PurePath
import re
import zipfile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError

from .i18n import translate

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000
MAX_COLUMNS = 50
MAX_CELLS = 100000
MAX_CELL_LENGTH = 10000
MAX_XLSX_UNCOMPRESSED = 50 * 1024 * 1024


def _string(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _table(rows, lang: str = "zh") -> pd.DataFrame:
    iterator = iter(rows)
    try:
        header = [_string(value) for value in next(iterator)]
    except StopIteration as exc:
        raise ValueError(translate("empty_table", lang)) from exc
    if not header or any(not name.strip() for name in header):
        raise ValueError(translate("empty_header", lang))
    if any(len(name) > MAX_CELL_LENGTH for name in header):
        raise ValueError(translate("long_header", lang, limit=MAX_CELL_LENGTH))
    if len(header) != len(set(header)):
        raise ValueError(translate("duplicate_header", lang))
    if len(header) > MAX_COLUMNS:
        raise ValueError(translate("too_many_columns", lang, limit=MAX_COLUMNS))
    body = []
    for row in iterator:
        if not row:
            continue
        values = [_string(value) for value in row]
        if len(values) != len(header):
            raise ValueError(translate("inconsistent_columns", lang))
        if any(len(value) > MAX_CELL_LENGTH for value in values):
            raise ValueError(translate("long_cell", lang, limit=MAX_CELL_LENGTH))
        body.append(values)
        if len(body) > MAX_ROWS or len(body) * len(header) > MAX_CELLS:
            raise ValueError(translate("too_many_cells", lang, rows=MAX_ROWS, cells=MAX_CELLS))
    if not body:
        raise ValueError(translate("header_only", lang))
    return pd.DataFrame(body, columns=header, dtype=object)


def read_dataset(data: bytes, filename: str, lang: str = "zh") -> pd.DataFrame:
    if not data:
        raise ValueError(translate("empty_file", lang))
    if len(data) > MAX_BYTES:
        raise ValueError(translate("file_too_large", lang))
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".csv":
        try:
            if data.startswith((b"\xff\xfe", b"\xfe\xff")):
                text = data.decode("utf-16")
            else:
                try:
                    text = data.decode("utf-8-sig")
                except UnicodeDecodeError:
                    text = data.decode("gb18030")
            if "\x00" in text:
                raise ValueError(translate("null_csv", lang))
            return _table(csv.reader(io.StringIO(text, newline=""), strict=True), lang)
        except (UnicodeError, csv.Error) as exc:
            raise ValueError(translate("invalid_csv", lang)) from exc
    if suffix != ".xlsx":
        raise ValueError(translate("unsupported_file", lang))
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 1000 or sum(info.file_size for info in archive.infolist()) > MAX_XLSX_UNCOMPRESSED:
                raise ValueError(translate("large_xlsx", lang))
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.worksheets[0]
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise ValueError(translate("wide_xlsx", lang, limit=MAX_COLUMNS))
            return _table(sheet.iter_rows(max_row=min(sheet.max_row or 1, MAX_ROWS + 2), values_only=True), lang)
        finally:
            workbook.close()
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ParseError, IndexError) as exc:
        raise ValueError(translate("invalid_xlsx", lang)) from exc


def _safe_cell(value) -> str:
    text = _string(value)
    if text.startswith(("\t", "\r", "\n")) or re.match(r"^\s*[=+\-@]", text):
        return "'" + text
    return text


def export_csv(frame: pd.DataFrame) -> bytes:
    """Prefix potential spreadsheet formulas, including formulas in headers."""
    safe = frame.map(_safe_cell)
    safe.columns = [_safe_cell(column) for column in frame.columns]
    return safe.to_csv(index=False, lineterminator="\n").encode("utf-8-sig")
